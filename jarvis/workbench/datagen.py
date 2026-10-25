"""Raw, messy source files for the five workbenches (Afterglow Labs, September 2026 close).

Everything is simulated, but shaped like what an accountant actually downloads: report title rows,
footers, mixed date formats, "$1,042.17" and "(95.00)" amounts, OCR noise, duplicate export rows,
inconsistent file names and vendor names. The agents must clean it before they can use it.

    python -m jarvis.workbench.datagen          -> data/workbench/<problem>/raw/...
"""
from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from jarvis import config

ROOT = config.ROOT / "data" / "workbench"


def _w(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text)


def _csv(path: Path, rows: list[list], bom: bool = False):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig" if bom else "utf-8") as fh:
        csv.writer(fh).writerows(rows)


# =============================================================================== 1. TIMING & ACCRUALS
def p1(base: Path):
    raw = base / "raw"
    _csv(raw / "bank" / "chase_x4471_export_2026-10.csv", [
        ["Details", "Posting Date", "Description", "Amount", "Type", "Balance", "Check or Slip #"],
        ["CREDIT", "10/01/2026", "INTEREST PAYMENT ", "$1,042.17", "MISC_CREDIT", "1,166,740.34", ""],
        ["CREDIT", "10/01/2026", "INTEREST PAYMENT ", "$1,042.17", "MISC_CREDIT", "1,166,740.34", ""],   # export glitch: duplicate
        ["DEBIT", "2026-10-02", "MONTHLY SERVICE FEE", "-95.00", "FEE_TRANSACTION", "1,166,645.34", ""],
        ["", "", "", "", "", "", ""],
        ["CREDIT", "10/05/26", "ORIG CO NAME:STRIPE           ORIG ID:1800948598 CO ENTRY DESCR:TRANSFER",
         "30115.22", "ACH_CREDIT", "1,196,760.56", ""],
        ["DEBIT", "10/05/2026", "ORIG CO NAME:RAMP             ORIG ID:1474827333 CO ENTRY DESCR:CARD PMT",
         "(49,306.71)", "ACH_DEBIT", "1,147,453.85", ""],
        ["DEBIT", "10/09/2026", "  ORIG CO NAME:GUSTO NET PAY ", "-95,112.40", "ACH_DEBIT", "1,052,341.45", ""],
    ])
    _w(raw / "bank" / "chase_interest_notice_2026-09.txt", """JPMORGAN CHASE BANK, N.A.            BUSINESS COMPLETE CHECKING  ...4471
AFTERGLOW LABS INC

  INTEREST ACCRUAL NOTICE
  ------------------------------------------------------------
  PERIOD           O9/01/2026 - O9/30/2026
  INTEREST ACCRUED                          $1,O42.17
  ANNUAL PERCENTAGE YIELD EARNED            0.7O%
  WILL BE CREDITED ON                       1O/01/2026
  ------------------------------------------------------------
  page 1 of 1     [scanned copy - OCR]
""")
    _csv(raw / "qbo" / "qbo_general_ledger_2026-09.csv", [
        ["Afterglow Labs, Inc."], ["General Ledger / Transaction Detail by Account"], ["September 2026"], [],
        ["Transaction Date", "Transaction Type", "Num", "Name", "Memo/Description", "Distribution Account", "Split", "Net Amount", "Debit", "Credit", "Balance"],
        ["1260 Interest Receivable", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "Beginning Balance", "", "", "", "", "", "994.95"],
        ["09/01/2026", "Deposit", "", "", "Interest received (Aug)", "1260", "1010 Cash - Chase Operating", "-994.95", "", "994.95", "0.00"],
        ["Total for 1260 Interest Receivable", "", "", "", "", "", "", "-994.95", "", "994.95", ""],
        ["2100 Accrued Expenses", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "Beginning Balance", "", "", "", "", "", "-13,450.00"],
        ["09/01/2026", "Journal Entry", "AJE-0801R", "Lumen Crew LLC", "Reverse Aug accrual", "2100", "5010 Crew & Production", "13,450.00", "13,450.00", "", "0.00"],
        ["Total for 2100 Accrued Expenses", "", "", "", "", "", "", "13,450.00", "13,450.00", "", ""],
        ["4900 Interest Income", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "Beginning Balance", "", "", "", "", "", "-7,049.77"],
        ["Total for 4900 Interest Income", "", "", "", "", "", "", "0.00", "", "", ""],
        ["5010 Crew & Production Costs", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "Beginning Balance", "", "", "", "", "", "47,712.20"],
        ["09/01/2026", "Journal Entry", "AJE-0801R", "Lumen Crew LLC", "Reverse Aug accrual", "5010", "2100 Accrued Expenses", "-13,450.00", "", "13,450.00", "34,262.20"],
        ["09/08/2026", "Bill", "LC-2026-08", "Lumen Crew LLC", "Lighting services 2026-08", "5010", "2010 Accounts Payable", "14,963.00", "14,963.00", "", "49,225.20"],
        ["09/09/2026", "Expense", "", "Damlicht Crew B.V.", "Crew fee Amsterdam 13 Sep", "5010", "1010 Cash - Chase Operating", "2,108.44", "2,108.44", "", "51,333.64"],
        ["Total for 5010 Crew & Production Costs", "", "", "", "", "", "", "3,621.44", "17,071.44", "13,450.00", ""],
        ["5020 Gear Rental", "", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "Beginning Balance", "", "", "", "", "", "1,680.00"],
        ["09/23/2026", "Expense", "", "Bay Rig Rentals", "Refund - 1 motor unused", "5020", "2050 Ramp Card Payable", "-420.00", "", "420.00", "1,260.00"],
        ["Total for 5020 Gear Rental", "", "", "", "", "", "", "-420.00", "", "420.00", ""],
        ["TOTAL", "", "", "", "", "", "", "", "", "", ""], [],
        ["Accrual Basis", "", "", "", "Tuesday, October 6, 2026 09:14 AM GMT-07:00", "", "", "", "", "", ""],
    ])
    _csv(raw / "qbo" / "ap_open_bills_2026-10-06.csv", [
        ["Vendor", "Bill Date", "Num", "Due Date", "Open Balance", "Memo"],
        ["Anthropic", "09/02/2026", "ANT-094471", "10/02/2026", "3,092.64", "API usage Sept"],
        ["lumen crew llc ", "9/8/2026", "LC-2026-08", "10/8/2026", "$14,963.00", "Lighting services 2026-08 (AUGUST)"],
        ["Amazon Web Services", "2026-09-02", "AWS-09-1188", "2026-10-02", "2,566.72", ""],
    ])
    _w(raw / "contracts" / "SOW_LumenCrew_2026_master.txt", """MASTER STATEMENT OF WORK — LIGHTING CREW
Vendor: Lumen Crew LLC (lumencrew.co)            Client: Afterglow Labs, Inc.
Term: June 1, 2026 – December 31, 2026
Rate: USD $95.00 per approved crew-hour. After-hours load-out (after 1:00 am) billed at 1.5x.
Billing: monthly in arrears, invoice by the 10th of the following month.
Signed: R. Ito (Lumen Crew)   S. Rivera (Afterglow)
""")
    _w(raw / "contracts" / "booking_NadiaKroll_FoxTheater_0925.txt", """BOOKING CONFIRMATION
Artist/Vendor: Nadia Kroll (VJ) — nadiakroll.studio
Event: Afterglow Live @ Fox Theater Oakland — September 25, 2026
Fee: $2,500.00 flat. Invoice after the show. Payment Net 15.
""")
    _w(raw / "contracts" / "SOW_Strobe_Portland_1010.txt", """SOW — Strobe & laser package
Vendor: Strobe & Co (strobeandco.com)
Event date: October 10, 2026 (Portland)
Fee: $3,000.00 flat
""")
    _csv(raw / "timesheets" / "LumenCrew_timesheet_SEPT (approved).csv", [
        ["Date", "Event", "Crew Member", "Hours", "Approved"],
        ["9/5/2026", "SF Mezzanine", "Rae Ito", "8.5h", "Y"],
        ["9/5/2026", "SF Mezzanine", "Tomas B.", "9", "Y"],
        ["09/05/2026", "SF Mezzanine", "Tomas B.", "9", "Y"],              # duplicate row
        ["9/5/2026", "SF Mezzanine", "Jules K", "9.5 hrs", "Y"],
        ["2026-09-19", "Oakland warehouse", "Rae Ito", "10", "yes"],
        ["2026-09-19", "Oakland warehouse", "Tomas B.", "eight", "yes"],
        ["2026-09-19", "Oakland warehouse", "Jules K", "9", "Y"],
        ["9/19/26", "Oakland warehouse", "Dee M.", "10", "Y"],
        ["9/19/26", "Oakland warehouse", "Sam P. (trainee)", "6", "N"],    # not approved
        ["9/25/2026", "Fox Theater", "Rae Ito", "9", "Y"],
        ["9/25/2026", "Fox Theater", "Tomas B.", "10", "Y"],
        ["9/25/2026", "Fox Theater", "Jules K", "9.5", "Y"],
        ["9/25/2026", "Fox Theater", "Dee M.", "9", "Y"],
        ["9/25/2026", "Fox Theater", "Kit L.", "10", "Y"],
        ["9/25/2026", "Fox Theater", "Ash N.", "9.5", "Y"],
        ["9/25/2026", "Fox Theater", "Bo C.", "9", "Y"],
        ["9/25/2026", "Fox Theater", "Rae Ito", "", "Y"],                  # blank hours
        ["", "", "TOTAL", "130", ""],
    ])
    _w(raw / "email" / "2026-09-29_Re_September_invoice_timing.eml", """From: Rae Ito <ops@lumencrew.co>
To: ap@afterglowlabs.com
Date: Tue, 29 Sep 2026 17:12:04 -0700
Subject: Re: September invoice timing

Hey! We're slammed with load-outs this week - September should land around 10k, give or take.
Invoice by the 10th like usual.
- Rae
""")
    _w(raw / "ops" / "run_of_show_FoxTheater_2026-09-25.txt", """RUN OF SHOW — Afterglow Live @ Fox Theater Oakland — Fri 25 Sep 2026
Lighting: Lumen Crew LLC (7 crew)   VJ: Nadia Kroll   VJ assist: Kai Mendes
Load-in 14:00 · Doors 20:00 · Show 21:30–01:00 · Load-out 01:00–03:30 (after-hours)
""")
    _w(raw / "invoices" / "BR-20871 BayRig.txt", """BAY RIG RENTALS INC. — INVOICE BR-20871
Invoice date: 10/03/2026            Terms: Net 30
Bill to: Afterglow Labs, Inc.
Service period: Sept 20 – Sept 27, 2026  (truss, 8x 1-ton motors, delivery)
TOTAL DUE: $3,600.00
""")
    _csv(raw / "history" / "accrual_history_2026.csv", [
        ["period", "vendor", "method", "accrued", "actual_invoice", "invoice_received", "notes"],
        ["2026-06", "Lumen Crew LLC", "rate x timesheet", "11,020.00", "12,250.00", "2026-07-08", "after-hours premium"],
        ["2026-07", "Lumen Crew LLC", "rate x timesheet", "12,825.00", "14,240.00", "2026-08-07", ""],
        ["2026-08", "Lumen Crew LLC", "rate x timesheet", "13,450.00", "14,963.00", "2026-09-08", ""],
        ["2026-07", "Nadia Kroll", "flat fee", "2,500.00", "2,250.00", "2026-08-30", "controller corrected: Nadia texted $2,250"],
        ["2026-02", "Nadia Kroll", "flat fee", "1,800.00", "1,800.00", "2026-04-02", "invoice 47 days late"],
    ])


# =============================================================================== 2. PACKAGE ROLL-FORWARD
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep"]
IS_ACCTS = [  # (num, name, section, monthly values Jan..Sep)  income positive, expense positive
    ("4010", "SaaS Subscription Revenue", "Income", [98200, 99800, 101400, 103900, 104700, 106100, 108300, 110200, 112750]),
    ("4020", "Live Show Revenue", "Income", [21000, 18500, 32000, 24500, 27800, 41000, 38500, 30200, 30575]),
    ("5010", "Crew & Production Costs", "Cost of Goods Sold", [6100, 5400, 9800, 7200, 8100, 13900, 15100, 14500, 3621]),
    ("5020", "Gear Rental", "Cost of Goods Sold", [1200, 900, 2100, 1600, 1500, 2800, 3100, 1680, 3180]),
    ("6100", "Salaries & Wages", "Expenses", [189300, 190100, 190800, 191200, 191900, 192400, 193100, 192800, 190757]),
    ("6410", "Software & AI Tools", "Expenses", [3900, 4100, 4050, 4200, 4300, 4450, 4500, 4610, 4790]),
    ("6420", "Advertising & Marketing", "Expenses", [600, 0, 900, 400, 1200, 0, 300, 450, 1850]),
    ("6430", "Travel", "Expenses", [2200, 1800, 3900, 2600, 2900, 4100, 4400, 3300, 3610]),
    ("6450", "Cloud Hosting", "Expenses", [2310, 2280, 2350, 2420, 2390, 2480, 2510, 2460, 2567]),
    ("6810", "Bank & FX Fees", "Expenses", [95, 95, 160, 95, 210, 170, 95, 150, 257]),
    ("4900", "Interest Income", "Other Income", [1010, 960, 1040, 1020, 1060, 990, 975, 995, 0]),
]
NEW_ACCT = ("6455", "AI Compute (GPU rental)", "Expenses", 6840)          # appears for the first time in September
RENAMED = ("6410", "Software & Subscriptions")                           # QBO name changed in September


def _pkg_numbers():
    """Monthly P&L + a balance sheet / cash flow that tie by construction."""
    ni = []
    for i in range(9):
        inc = sum(a[3][i] for a in IS_ACCTS if a[2] in ("Income", "Other Income"))
        exp = sum(a[3][i] for a in IS_ACCTS if a[2] in ("Cost of Goods Sold", "Expenses")) + (NEW_ACCT[3] if i == 8 else 0)
        ni.append(inc - exp)
    bs = {"1260": [1010, 960, 1040, 1020, 1060, 990, 975, 995, 0], "1310": [0, 0, 0, 0, 2860, 2860, 0, 0, 3452],
          "2010": [7800, 6900, 9100, 8800, 9600, 12400, 14100, 16700, 20621], "2050": [9400, 8800, 10900, 9700, 11200, 12800, 13400, 4931, 11131],
          "2100": [0, 0, 0, 2700, 0, 11020, 15075, 13450, 0]}
    ytd, cash = [], []
    open_eq, open_cash = 1_850_000, 1_850_000
    run = 0
    for i in range(9):
        run += ni[i]
        ytd.append(run)
        liab = bs["2010"][i] + bs["2050"][i] + bs["2100"][i]
        cash.append(open_eq + run + liab - bs["1260"][i] - bs["1310"][i])
    return ni, ytd, bs, cash, open_eq


def p2(base: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    raw = base / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    ni, ytd, bs, cash, open_eq = _pkg_numbers()
    hdr = Font(bold=True)
    band = PatternFill("solid", fgColor="E8E3FF")

    # ---------- last month's package (through August), as the team keeps it
    wb = Workbook()
    ws = wb.active
    ws.title = "IS"
    ws.append(["Afterglow Labs, Inc. — Income Statement (monthly)"])
    ws.append(["Acct #", "Account"] + [f"{m} 2026" for m in MONTHS[:8]] + ["YTD"])
    for c in ws[2]:
        c.font, c.fill = hdr, band
    for section in ("Income", "Cost of Goods Sold", "Expenses", "Other Income"):
        ws.append(["", section.upper()])
        first = ws.max_row + 1
        for a in [a for a in IS_ACCTS if a[2] == section]:
            r = ws.max_row + 1
            ws.append([a[0], a[1]] + a[3][:8] + [f"=SUM(C{r}:J{r})"])
        last = ws.max_row
        ws.append(["", f"Total {section}"] + [f"=SUM({col}{first}:{col}{last})" for col in "CDEFGHIJK"])
        ws[ws.max_row][1].font = hdr
    ws.append(["", "NET INCOME"] + [ni[i] for i in range(8)] + [sum(ni[:8])])
    ws[ws.max_row][1].font = hdr
    ws.column_dimensions["B"].width = 34
    bsws = wb.create_sheet("BS")
    bsws.append(["Afterglow Labs, Inc. — Balance Sheet (month end)"])
    bsws.append(["Acct #", "Account"] + [f"{m} 2026" for m in MONTHS[:8]])
    for c in bsws[2]:
        c.font, c.fill = hdr, band
    rows = [("1010", "Cash - Chase Operating", cash), ("1260", "Interest Receivable", bs["1260"]),
            ("1310", "Prepaid Tour Costs", bs["1310"]), ("", "Total Assets", None),
            ("2010", "Accounts Payable", bs["2010"]), ("2050", "Ramp Card Payable", bs["2050"]),
            ("2100", "Accrued Expenses", bs["2100"]), ("3000", "Opening Balance Equity", [open_eq] * 9),
            ("3900", "Net Income (YTD)", ytd), ("", "Total Liabilities & Equity", None)]
    for num, name, vals in rows:
        if vals is None:
            r = bsws.max_row
            span = (r - 2, r) if name == "Total Assets" else (r - 4, r)
            bsws.append(["", name] + [f"=SUM({c}{span[0]}:{c}{span[1]})" for c in "CDEFGHIJ"])
            bsws[bsws.max_row][1].font = hdr
        else:
            bsws.append([num, name] + vals[:8])
    bsws.column_dimensions["B"].width = 34
    cf = wb.create_sheet("CF")
    cf.append(["Afterglow Labs, Inc. — Statement of Cash Flows (monthly, indirect)"])
    cf.append(["", "Line"] + [f"{m} 2026" for m in MONTHS[:8]])
    for c in cf[2]:
        c.font, c.fill = hdr, band
    def d(series, i, sign):
        prev = series[i - 1] if i else (1010 if series is bs["1260"] else (9400 if series is bs["2050"] else 0))
        return sign * (series[i] - prev)
    cf_rows = [("Net Income", ni), ("Change in Interest Receivable", [d(bs["1260"], i, -1) for i in range(9)]),
               ("Change in Prepaid Tour Costs", [d(bs["1310"], i, -1) for i in range(9)]),
               ("Change in Accounts Payable", [d(bs["2010"], i, 1) for i in range(9)]),
               ("Change in Ramp Card Payable", [d(bs["2050"], i, 1) for i in range(9)]),
               ("Change in Accrued Expenses", [d(bs["2100"], i, 1) for i in range(9)])]
    for name, vals in cf_rows:
        cf.append(["", name] + vals[:8])
    notes = wb.create_sheet("Flux comments")
    notes.append(["Account", "Month", "Change", "Comment", "Preparer"])
    notes.append(["5010 Crew & Production Costs", "Aug 2026", -600, "Fewer shows in August; Lumen invoice timing.", "Dev Patel"])
    wb.save(raw / "Afterglow_FinancialPackage_2026-08_FINAL.xlsx")

    # ---------- QBO report exports for September (title rows, indents, totals, footer, text numbers)
    def qbo_export(path: Path, title: str, period: str, lines: list[tuple[str, object]], footer_time="09:14 AM"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Afterglow Labs, Inc."])
        ws.append([title])
        ws.append([period])
        ws.append([])
        ws.append(["", "Total"])
        for label, val in lines:
            ws.append([label, val])
            indent = len(label) - len(label.lstrip(" "))
            if indent:
                ws.cell(ws.max_row, 1).alignment = Alignment(indent=indent // 3)
        ws.append([])
        ws.append([f"Accrual Basis Tuesday, October 6, 2026 {footer_time} GMT-07:00"])
        wb.save(path)

    def money_text(v):  # QBO exports sometimes land as text with separators
        return f"{v:,.2f}"

    def pl_lines(extra_5010=0):
        out = []
        for section, label in (("Income", "Income"), ("Cost of Goods Sold", "Cost of Goods Sold"), ("Expenses", "Expenses")):
            out.append((label, None))
            tot = 0
            for a in [a for a in IS_ACCTS if a[2] == section]:
                v = a[3][8] + (extra_5010 if a[0] == "5010" else 0)
                name = RENAMED[1] if a[0] == RENAMED[0] else a[1]
                out.append((f"   {a[0]} {name}", money_text(v) if a[0] in ("6100", "4010") else v))
                tot += v
            if section == "Expenses":
                out.append((f"   {NEW_ACCT[0]} {NEW_ACCT[1]}", NEW_ACCT[3]))
                tot += NEW_ACCT[3]
            out.append((f"Total {label}", tot))
            if section == "Cost of Goods Sold":
                inc = sum(a[3][8] for a in IS_ACCTS if a[2] == "Income")
                out.append(("Gross Profit", inc - tot))
        oi = [a for a in IS_ACCTS if a[2] == "Other Income"]
        out.append(("Other Income", None))
        for a in oi:
            out.append((f"   {a[0]} {a[1]}", a[3][8]))
        out.append(("Total Other Income", sum(a[3][8] for a in oi)))
        out.append(("Net Income", ni[8] - extra_5010))
        return out

    def bs_lines(extra_accrued=0):
        assets = [("1010", "Cash - Chase Operating", cash[8]), ("1260", "Interest Receivable", bs["1260"][8]),
                  ("1310", "Prepaid Tour Costs", bs["1310"][8])]
        liab = [("2010", "Accounts Payable", bs["2010"][8]), ("2050", "Ramp Card Payable", bs["2050"][8]),
                ("2100", "Accrued Expenses", bs["2100"][8] + extra_accrued)]
        eq = [("3000", "Opening Balance Equity", open_eq), ("3900", "Net Income", ytd[8] - extra_accrued)]
        out = [("ASSETS", None), ("   Current Assets", None)]
        out += [(f"      {n} {nm}", v) for n, nm, v in assets]
        out += [("   Total Current Assets", sum(v for *_, v in assets)), ("TOTAL ASSETS", sum(v for *_, v in assets)),
                ("LIABILITIES AND EQUITY", None), ("   Liabilities", None)]
        out += [(f"      {n} {nm}", v) for n, nm, v in liab]
        out += [("   Total Liabilities", sum(v for *_, v in liab)), ("   Equity", None)]
        out += [(f"      {n} {nm}", v) for n, nm, v in eq]
        out += [("   Total Equity", sum(v for *_, v in eq)),
                ("TOTAL LIABILITIES AND EQUITY", sum(v for *_, v in liab) + sum(v for *_, v in eq))]
        return out

    def cf_lines(extra=0):
        rows = [("OPERATING ACTIVITIES", None), ("   Net Income", ni[8] - extra),
                ("   Adjustments to reconcile Net Income to Net Cash provided by operations:", None),
                ("      Interest Receivable", -(bs["1260"][8] - bs["1260"][7])),
                ("      Prepaid Tour Costs", -(bs["1310"][8] - bs["1310"][7])),
                ("      Accounts Payable", bs["2010"][8] - bs["2010"][7]),
                ("      Ramp Card Payable", bs["2050"][8] - bs["2050"][7]),
                ("      Accrued Expenses", bs["2100"][8] - bs["2100"][7] + extra)]
        op = sum(v for _, v in rows if isinstance(v, int))
        rows += [("Net cash provided by operating activities", op), ("NET CASH INCREASE FOR PERIOD", op),
                 ("Cash at beginning of period", cash[7]), ("CASH AT END OF PERIOD", cash[7] + op)]
        return rows

    qbo_export(raw / "ProfitAndLoss (Sept).xlsx", "Profit and Loss", "September 2026", pl_lines())
    qbo_export(raw / "BalanceSheet.xlsx", "Balance Sheet", "As of September 30, 2026", bs_lines())
    qbo_export(raw / "Statement of Cash Flows.xlsx", "Statement of Cash Flows", "September 2026", cf_lines())
    late = base / "late_change"
    late.mkdir(parents=True, exist_ok=True)
    qbo_export(late / "ProfitAndLoss (Sept) (1).xlsx", "Profit and Loss", "September 2026", pl_lines(12350), "04:51 PM")
    qbo_export(late / "BalanceSheet (1).xlsx", "Balance Sheet", "As of September 30, 2026", bs_lines(12350), "04:51 PM")
    qbo_export(late / "Statement of Cash Flows (1).xlsx", "Statement of Cash Flows", "September 2026", cf_lines(12350), "04:52 PM")
    _w(late / "README_late_change.txt", "Controller posted AJE-0930-04 at 4:47 PM: accrue Lumen Crew September "
       "lighting $12,350.00 (Dr 5010 / Cr 2100). Package must be refreshed before the board pack goes out.\n")
    _csv(raw / "gl_detail_top_movers_2026-09.csv", [
        ["Account", "Date", "Name", "Memo", "Amount"],
        ["5010 Crew & Production Costs", "09/01/2026", "Lumen Crew LLC", "Reverse Aug accrual", "-13,450.00"],
        ["5010 Crew & Production Costs", "09/08/2026", "Lumen Crew LLC", "Aug invoice LC-2026-08", "14,963.00"],
        ["5010 Crew & Production Costs", "09/09/2026", "Damlicht Crew B.V.", "Crew fee Amsterdam 13 Sep", "2,108.44"],
        ["6455 AI Compute (GPU rental)", "09/14/2026", "Lambda Labs", "GPU cluster for real-time visuals beta", "6,840.00"],
        ["6420 Advertising & Marketing", "09/08/2026", "Google Ads", "Fall tour promo", "1,250.00"],
        ["6420 Advertising & Marketing", "09/12/2026", "Lumen Studios", "Promo photo shoot", "600.00"],
        ["4900 Interest Income", "", "", "No September interest recorded (credited 10/01 - see timing workbench)", "0.00"],
        ["4020 Live Show Revenue", "09/23/2026", "Desert Hearts LLC", "Wire - festival settlement", "30,574.75"],
        ["5020 Gear Rental", "09/20/2026", "Bay Rig Rentals", "Rigging Sept 20-27 (accrual pending)", "0.00"],
    ])


# =============================================================================== 3. BANK TRANSACTIONS
def p3(base: Path):
    raw = base / "raw"
    _csv(raw / "qbo_bank_feed_for_review_2026-09.csv", [
        ["Date", "Description", "Payee", "Category or Match", "Spent", "Received"],
        ["09/03/2026", "ORIG CO NAME:STRIPE TRANSFER ST-PO8812", "Stripe", "Match: Deposit 09/03 $30,615.46", "", "30,615.46"],
        ["09/09/2026", "PAYPAL *INST XFER 16F1203617 WEB ID: PAYPALSI77", "PayPal", "Software & AI Tools", "2,262.21", ""],
        ["09/11/2026", "SQ *MOONLIGHT GEAR CO  Oakland CA", "", "Uncategorized Expense", "742.18", ""],
        ["9/15/2026", "ORIG CO NAME:GUSTO NET PAY", "Gusto", "Match: Payroll 09/15", "94,764.63", ""],
        ["09/17/2026", "AMZN Mktp US*2K4LM19Q2 Amzn.com/bill WA", "Amazon", "Office & Misc", "214.37", ""],
        ["09/18/2026", "PAYPAL *INST XFER DCE4343478 WEB ID: PAYPALSI77", "PayPal", "Software & AI Tools", "3,912.44", ""],
        ["09/24/2026", "PAYPAL *INST XFER DDC0867643 WEB ID: PAYPALSI77", "PayPal", "Software & AI Tools", "1,000.00", ""],
        ["09/28/2026", "Zelle payment to R Lopez JPM99c92a10", "", "Uncategorized Expense", "350.00", ""],
        ["09/30/2026", "MONTHLY SERVICE FEE", "JPMorgan Chase", "Bank Charges", "95.00", ""],
    ])
    _csv(raw / "paypal_activity_download_2026-09.csv", [
        ["Date", "Time", "TimeZone", "Name", "Type", "Status", "Currency", "Gross", "Fee", "Net", "From Email Address",
         "To Email Address", "Transaction ID", "Item Title", "Reference Txn ID", "Balance Impact"],
        ["09/09/2026", "09:28:00", "America/Los_Angeles", "", "Bank Deposit to PP Account", "Completed", "USD", "2,262.21", "0.00", "2,262.21", "", "", "16F12036174027107", "", "", "Credit"],
        ["09/09/2026", "10:05:01", "America/Los_Angeles", "Damlicht Crew B.V.", "General Payment", "Completed", "EUR", "-1,800.00", "0.00", "-1,800.00", "ap@afterglowlabs.com", "billing@damlicht.nl", "8AF529710FA489324", "Crew fee Amsterdam 13 Sep", "", "Debit"],
        ["09/09/2026", "10:05:01", "America/Los_Angeles", "", "General Currency Conversion", "Completed", "USD", "-2,237.22", "0.00", "-2,237.22", "", "", "33E727668ED704645", "", "8AF529710FA489324", "Debit"],
        ["09/09/2026", "10:05:01", "America/Los_Angeles", "", "General Currency Conversion", "Completed", "EUR", "1,800.00", "0.00", "1,800.00", "", "", "48386339566115370", "", "8AF529710FA489324", "Credit"],
        ["09/09/2026", "18:29:27", "America/Los_Angeles", "Splice.com", "PreApproved Payment Bill User Payment", "Completed", "USD", "-24.99", "0.00", "-24.99", "ap@afterglowlabs.com", "billing@splice.com", "BB999784197763884", "Splice Sounds subscription", "", "Debit"],
        ["09/18/2026", "09:22:15", "America/Los_Angeles", "Kraftwerk Rigging GmbH", "General Payment", "Completed", "EUR", "-2,950.00", "0.00", "-2,950.00", "ap@afterglowlabs.com", "billing@kraftwerk-rigging.de", "6B5301917A2670494", "Deposit rigging Berlin 14 Nov", "", "Debit"],
        ["09/18/2026", "09:22:15", "America/Los_Angeles", "", "General Currency Conversion", "Completed", "USD", "-3,599.00", "0.00", "-3,599.00", "", "", "F596647782A987131", "", "6B5301917A2670494", "Debit"],
        ["09/18/2026", "09:22:15", "America/Los_Angeles", "", "General Currency Conversion", "Completed", "EUR", "2,950.00", "0.00", "2,950.00", "", "", "802659060E0457522", "", "6B5301917A2670494", "Credit"],
        ["09/18/2026", "09:58:35", "America/Los_Angeles", "Valve Corporation", "Express Checkout Payment", "Completed", "USD", "-113.44", "0.00", "-113.44", "j.ortiz@afterglowlabs.com", "noreply@steampowered.com", "C15807984D0797968", "Steam purchase", "", "Debit"],
        ["09/18/2026", "10:28:29", "America/Los_Angeles", "Anthropic, PBC", "PreApproved Payment Bill User Payment", "Completed", "USD", "-200.00", "0.00", "-200.00", "ap@afterglowlabs.com", "billing@anthropic.com", "8971295127F968416", "Claude Max subscription", "", "Debit"],
        ["09/18/2026", "14:48:00", "America/Los_Angeles", "", "Bank Deposit to PP Account", "Completed", "USD", "3,912.44", "0.00", "3,912.44", "", "", "DCE43434786208450", "", "6B5301917A2670494", "Credit"],
        ["09/24/2026", "11:02:00", "America/Los_Angeles", "", "Bank Deposit to PP Account", "Completed", "USD", "1,000.00", "0.00", "1,000.00", "", "", "DDC08676432295501", "", "", "Credit"],
        ["09/26/2026", "16:40:12", "America/Los_Angeles", "Kai Mendes", "General Payment", "Completed", "USD", "-640.00", "0.00", "-640.00", "ap@afterglowlabs.com", "kai@kaimendes.art", "4K2207731MN090317", "VJ assist - Oakland 25 Sep", "", "Debit"],
    ], bom=True)
    _w(raw / "receipts" / "Square_receipt_MoonlightGear_0911.txt", """MOONLIGHT GEAR CO · Oakland CA · via Square
09/11/2026 2:14 PM   Receipt #2X7Q
2x DMX512 cable 50ft ........ 118.00
1x Haze fluid 4L ............  89.00
1x Gaffer tape (12 rolls) ...  96.00
4x LED par replacement lens .. 380.00
Subtotal 683.00 · Tax 59.18 · TOTAL $742.18   Card ....0091 (Ramp? no — debit 4471)
Customer: Sam Rivera (Afterglow Labs) — "for Fox Theater 9/25"
""")
    _w(raw / "receipts" / "Amazon_order_112-4471029-2K4LM19.txt", """Amazon.com order 112-4471029-2K4LM19 · placed 09/16/2026 · charged 09/17/2026
Ship to: Afterglow Labs, Oakland CA
1x Brother TN-760 toner ............ 62.99
2x Rocketbook Core notebooks ........ 49.98
1x Elgato Stream Deck MK.2 ........ 149.99   (rebate -48.59)
Order total $214.37 · Payment: Chase debit ....4471
""")
    _csv(raw / "reference" / "ecb_eur_usd_2026-09.csv", [["date", "eur_usd"], ["2026-09-09", "1.1890"], ["2026-09-18", "1.1700"],
                                                          ["2026-09-24", "1.1712"], ["2026-09-26", "1.1695"]])
    _w(raw / "docs" / "SOW_Kraftwerk_Berlin_2026-11-14.txt", """STATEMENT OF WORK — Kraftwerk Rigging GmbH (kraftwerk-rigging.de)
Event: Afterglow Live @ Kraftwerk Berlin — 14 November 2026
Total EUR 9,800.00 · Deposit EUR 2,950.00 due 18 September 2026 via PayPal · balance 7 days after show
""")
    _w(raw / "docs" / "RunOfShow_Amsterdam_2026-09-13.txt", "RUN OF SHOW — Paradiso Amsterdam — 13 Sep 2026 — Crew: Damlicht Crew B.V. (6)\n")
    _w(raw / "docs" / "RunOfShow_Oakland_2026-09-25.txt", "RUN OF SHOW — Fox Theater Oakland — 25 Sep 2026 — VJ assist: Kai Mendes · Stagehand (Zelle): R. Lopez\n")
    _w(raw / "docs" / "policy_P-0317_personal_purchases.txt", "Controller ruling P-0317 (Mar 2026): personal purchases on company cards/PayPal "
       "go to 1250 Employee Receivable regardless of amount; notify the employee.\n")


# =============================================================================== 4. EVIDENCE & FOLDERS
def _pdf(path: Path, lines: list[str]):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setTitle(path.stem)
    y = 740
    for i, ln in enumerate(lines):
        c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 13 if i == 0 else 10.5)
        c.drawString(60, y, ln)
        y -= 20 if i == 0 else 16
    c.save()


def p4(base: Path):
    dl = base / "raw" / "Downloads"
    if dl.exists():
        shutil.rmtree(dl)
    op_sep = ["JPMorgan Chase Bank, N.A. — Business Complete Checking", "Account: AFTERGLOW LABS INC  ...4471",
              "Statement period: September 1, 2026 through September 30, 2026",
              "Beginning balance: $1,236,375.83", "Deposits and additions: $150,904.86",
              "Withdrawals and debits: -$221,582.52", "Ending balance: $1,165,698.17"]
    _pdf(dl / "Statement (3).pdf", op_sep)
    shutil.copy(dl / "Statement (3).pdf", dl / "eStatement_2026-09-30 (1).pdf")           # exact duplicate download
    _pdf(dl / "Statement (2).pdf", ["JPMorgan Chase Bank, N.A. — Business Complete Checking", "Account: AFTERGLOW LABS INC  ...4471",
                                     "Statement period: August 1, 2026 through August 31, 2026",
                                     "Beginning balance: $1,301,442.10", "Ending balance: $1,236,375.83"])
    _pdf(dl / "chase sav sept.pdf", ["JPMorgan Chase Bank, N.A. — Business Premier Savings", "Account: AFTERGLOW LABS INC  ...8820",
                                      "Statement period: September 1, 2026 through September 30, 2026",
                                      "Beginning balance: $400,000.00", "Interest paid: $233.10", "Ending balance: $400,233.10"])
    _pdf(dl / "ramp_statement_sep.pdf", ["Ramp — Corporate Card Statement", "Business: Afterglow Labs, Inc.   Card program ...0091",
                                          "Statement period: Sep 1, 2026 – Sep 30, 2026", "New balance: $11,131.41",
                                          "Payment due: Oct 5, 2026"])
    _pdf(dl / "QBO_Reconciliation_Report_Operating.pdf", ["Afterglow Labs, Inc. — Reconciliation Report",
                                                           "Account: 1010 Cash - Chase Operating (...4471)",
                                                           "Statement ending date: 09/30/2026", "Statement ending balance: 1,165,698.17",
                                                           "Cleared balance: 1,165,698.17   Difference: 0.00",
                                                           "Reconciled on: 10/06/2026 by Dev Patel"])
    _pdf(dl / "recon_final_FINAL_v2.pdf", ["Afterglow Labs, Inc. — Reconciliation Report", "Account: 1020 Cash - Chase Savings (...8820)",
                                            "Statement ending date: 09/30/2026", "Statement ending balance: 400,233.10",
                                            "Cleared balance: 400,233.10   Difference: 0.00", "Reconciled on: 10/06/2026 by Dev Patel"])
    _pdf(dl / "Reconciliation (1).pdf", ["Afterglow Labs, Inc. — Reconciliation Report", "Account: 2050 Ramp Card Payable (...0091)",
                                          "Statement ending date: 08/31/2026", "Statement ending balance: 4,930.67",
                                          "Cleared balance: 4,930.67   Difference: 0.00", "Reconciled on: 09/07/2026 by Dev Patel"])
    _pdf(dl / "Q4 budget draft v3.pdf", ["Afterglow Labs — Q4 budget (DRAFT)", "Not a close document."])
    _csv(base / "raw" / "account_inventory.csv", [
        ["Account", "Institution", "Last4", "QBO Account", "Requires"],
        ["Chase Operating", "JPMorgan Chase", "4471", "1010 Cash - Chase Operating", "statement;reconciliation"],
        ["Chase Savings", "JPMorgan Chase", "8820", "1020 Cash - Chase Savings", "statement;reconciliation"],
        ["Ramp Card", "Ramp", "0091", "2050 Ramp Card Payable", "statement;reconciliation"],
        ["PayPal Business", "PayPal", "PP01", "1030 PayPal Balance", "statement;reconciliation"],
    ])
    _w(base / "raw" / "drive_tree_existing.txt", """Afterglow Labs/Finance/Monthly Close/2026/2026-08/01 - Bank and Cash/01 - Statements/Chase Operating/2026-08_BankStatement_Chase-Operating-4471_2026-08-31.pdf
Afterglow Labs/Finance/Monthly Close/2026/2026-08/01 - Bank and Cash/01 - Statements/Chase Savings/2026-08_BankStatement_Chase-Savings-8820_2026-08-31.pdf
Afterglow Labs/Finance/Monthly Close/2026/2026-08/01 - Bank and Cash/01 - Statements/Ramp Card/2026-08_CardStatement_Ramp-0091_2026-08-31.pdf
Afterglow Labs/Finance/Monthly Close/2026/2026-08/01 - Bank and Cash/02 - Reconciliations/Chase Operating/2026-08_BankReconciliation_QBO-Chase-Operating-4471_2026-08-31.pdf
Afterglow Labs/Finance/Monthly Close/2026/2026-08/01 - Bank and Cash/02 - Reconciliations/Ramp Card/2026-08_BankReconciliation_QBO-Ramp-0091_2026-08-31.pdf
""")


# =============================================================================== 5. DUPLICATE VENDORS
def p5(base: Path):
    raw = base / "raw"
    _csv(raw / "qbo_vendor_list_export_2026-10-06.csv", [
        ["Vendor", "Company Name", "Email", "Tax ID", "Open Balance", "Created", "Created By"],
        ["Anthropic", "Anthropic, PBC", "billing@anthropic.com", "84-1234401", "3,092.64", "2025-03-11", "Dev Patel"],
        ["ANTHROPIC PBC ", "", "invoices@mail.anthropic.com", "", "0.00", "2026-06-02", "Bank rule"],
        ["Claude", "", "", "", "0.00", "2026-09-02", "Ramp sync"],
        ["Anysphere", "Anysphere, Inc.", "ap@cursor.com", "88-2214117", "480.00", "2025-11-20", "Dev Patel"],
        ["Cursor", "", "", "", "0.00", "2026-09-04", "Ramp sync"],
        ["Google", "Google LLC", "collections@google.com", "77-0493581", "0.00", "2024-08-01", "Maya Chen"],
        ["Google Ads", "", "ads-noreply@google.com", "", "0.00", "2026-09-08", "Ramp sync"],
        ["Google Workspace", "Google LLC", "workspace-noreply@google.com", "77-0493581", "0.00", "2025-01-15", "Ramp sync"],
        ["Amazon Web Services", "Amazon Web Services, Inc.", "aws-receivables@amazon.com", "91-1646860", "2,566.72", "2024-08-01", "Maya Chen"],
        ["AWS", "", "no-reply@aws.amazon.com", "", "0.00", "2026-02-12", "Ramp sync"],
        ["Splice", "Splice.com, Inc.", "billing@splice.com", "46-3317982", "0.00", "2025-05-02", "Dev Patel"],
        ["Splice.com", "", "billing@splice.com", "", "0.00", "2026-09-09", "PayPal rule"],
        ["Lumen Crew LLC", "Lumen Crew LLC", "ops@lumencrew.co", "83-4419210", "14,963.00", "2026-05-28", "Dev Patel"],
        ["Lumen Studios", "Lumen Studios SF LLC", "bookings@lumenstudios.com", "27-9031551", "0.00", "2026-09-12", "Ramp sync"],
        ["Delta Air Lines", "Delta Air Lines, Inc.", "", "58-0218548", "0.00", "2024-08-01", "Maya Chen"],
        ["Uber", "Uber Technologies, Inc.", "", "45-2647441", "0.00", "2024-08-01", "Maya Chen"],
    ])
    _csv(raw / "ramp_vendors_2026-10-06.csv", [
        ["Ramp Vendor", "Merchant Descriptor", "Receipt Issuer", "Receipt Domain", "Suggested Action", "ERP Vendor"],
        ["Claude", "CLAUDE.AI SUBSCRIPTION", "Anthropic, PBC", "anthropic.com", "Create new vendor 'Claude'", ""],
        ["Cursor", "CURSOR AI POWERED IDE", "Anysphere, Inc.", "cursor.com", "Create new vendor 'Cursor'", ""],
        ["Google Ads", "GOOGLE *ADS4471029388", "Google LLC", "google.com", "Create new vendor 'Google Ads'", ""],
        ["Lumen Studios", "LUMEN STUDIOS SF", "Lumen Studios SF LLC", "lumenstudios.com", "Create new vendor 'Lumen Studios'", ""],
        ["AWS", "AWS EMEA MARKETPLACE", "Amazon Web Services, Inc.", "aws.amazon.com", "Create new vendor 'AWS'", ""],
    ])
    _csv(raw / "vendor_transactions_2026.csv", [
        ["Source", "Date", "Vendor (as recorded)", "Amount", "Account", "Memo"],
        ["QBO bill", "2026-07-02", "Anthropic", "2,940.10", "6410", "API usage Jul"],
        ["QBO bill", "2026-08-03", "Anthropic", "3,212.55", "6410", "API usage Aug"],
        ["QBO bill", "2026-09-02", "Anthropic", "3,092.64", "6410", "API usage Sep"],
        ["Bank rule", "2026-06-02", "ANTHROPIC PBC ", "1,200.00", "6410", "Prepaid credits"],
        ["Ramp", "2026-09-02", "Claude", "150.00", "6410", "Claude Pro team seats"],
        ["PayPal", "2026-09-18", "Claude", "200.00", "6410", "Claude Max subscription"],
        ["QBO bill", "2026-09-06", "Anysphere", "480.00", "6410", "Cursor Business Sept"],
        ["Ramp", "2026-09-04", "Cursor", "480.00", "6410", "Cursor Business Sept"],
        ["Ramp", "2026-09-03", "Google Workspace", "756.00", "6410", "Workspace seats"],
        ["Ramp", "2026-09-08", "Google Ads", "1,250.00", "6420", "Fall tour promo"],
        ["QBO bill", "2026-01-15", "Google", "4,000.00", "6420", "Annual ads credit"],
        ["QBO bill", "2026-09-02", "Amazon Web Services", "2,566.72", "6450", "AWS Sept"],
        ["Ramp", "2026-02-12", "AWS", "312.00", "6450", "Marketplace AMI"],
        ["PayPal", "2026-08-19", "Splice", "24.99", "6410", "Splice Sounds"],
        ["PayPal", "2026-09-09", "Splice.com", "24.99", "6410", "Splice Sounds"],
        ["QBO bill", "2026-09-08", "Lumen Crew LLC", "14,963.00", "5010", "Lighting Aug"],
        ["Ramp", "2026-09-12", "Lumen Studios", "600.00", "6420", "Promo photo shoot studio rental"],
    ])
    _csv(raw / "w9_registry.csv", [
        ["Legal Name", "TIN", "Domain(s)", "Products / DBA", "W-9 on file"],
        ["Anthropic, PBC", "84-1234401", "anthropic.com;mail.anthropic.com", "Claude;Claude Pro;Claude Max;Claude API", "2025-03-11"],
        ["Anysphere, Inc.", "88-2214117", "cursor.com;anysphere.inc", "Cursor;Cursor Business", "2025-11-20"],
        ["Google LLC", "77-0493581", "google.com", "Google Ads;Google Workspace;Google Cloud", "2024-08-01"],
        ["Amazon Web Services, Inc.", "91-1646860", "amazon.com;aws.amazon.com", "AWS;AWS Marketplace", "2024-08-01"],
        ["Splice.com, Inc.", "46-3317982", "splice.com", "Splice;Splice Sounds", "2025-05-02"],
        ["Lumen Crew LLC", "83-4419210", "lumencrew.co", "lighting crew", "2026-05-28"],
        ["Lumen Studios SF LLC", "27-9031551", "lumenstudios.com", "photo studio rental", "2026-09-12"],
    ])


PROBLEMS = {"p1_timing": p1, "p2_rollforward": p2, "p3_bank": p3, "p4_evidence": p4, "p5_vendors": p5}


def generate(root: Path = ROOT) -> Path:
    for name, fn in PROBLEMS.items():
        base = root / name
        if base.exists():
            shutil.rmtree(base)
        base.mkdir(parents=True)
        fn(base)
    return root


if __name__ == "__main__":
    out = generate()
    for p in sorted(out.rglob("*")):
        if p.is_file():
            print(p.relative_to(out), p.stat().st_size)
