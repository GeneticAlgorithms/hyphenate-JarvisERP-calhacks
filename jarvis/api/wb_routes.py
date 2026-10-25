"""HTTP routes for the five agent workbenches."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from fastapi import APIRouter, Body, HTTPException
from fastapi.responses import FileResponse, JSONResponse

from jarvis.work import events
from jarvis.work.roles import ROLES
from jarvis.work.transition import Denied, get_item, transition
from jarvis.workbench.core import DATA, OUT, Workbench
from jarvis.workbench.scenarios import SCENARIOS
from jarvis.work.roles import actor

router = APIRouter()
WB: dict = {"wb": None, "speed": 0.18}


def W() -> Workbench:
    if WB["wb"] is None:
        WB["wb"] = Workbench(speed=WB["speed"])
    return WB["wb"]


@router.get("/api/wb/problems")
def problems():
    w = W()
    out = []
    for key, sc in SCENARIOS.items():
        base = DATA / key
        files = [str(p.relative_to(base)) for p in sorted(base.rglob("*")) if p.is_file()]
        agents = [{"id": a, "title": w.agents[a]["title"], "stages": w.agents[a]["stages"]} for a in sc["agents"]]
        out.append({"id": key, **sc, "files": files, "agent_specs": agents})
    from jarvis.world.bible import COA
    accts = {k: v[1] for k, v in COA.items()}
    return {"problems": out, "accounts": accts, "roles": {k: {r: sorted(a) for r, a in v.items()} for k, v in ROLES.items()}}


@router.get("/api/wb/raw/{problem}/{path:path}")
def raw(problem: str, path: str):
    p = (DATA / problem / path).resolve()
    if not str(p).startswith(str((DATA / problem).resolve())) or not p.exists():
        raise HTTPException(404)
    if p.suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p)
        sheets = {ws.title: [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
                  for ws in wb.worksheets}
        return {"kind": "xlsx", "sheets": sheets, "name": p.name}
    if p.suffix == ".pdf":
        from pypdf import PdfReader
        return {"kind": "pdf", "text": "\n".join(pg.extract_text() or "" for pg in PdfReader(str(p)).pages), "name": p.name,
                "url": f"/api/wb/rawfile/{problem}/{path}"}
    text = p.read_text(encoding="utf-8", errors="replace")
    if p.suffix == ".csv":
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        return {"kind": "csv", "rows": rows, "text": text, "name": p.name, "bom": text.startswith("﻿")}
    return {"kind": "text", "text": text, "name": p.name}


@router.get("/api/wb/rawfile/{problem}/{path:path}")
def rawfile(problem: str, path: str):
    p = (DATA / problem / path).resolve()
    if not str(p).startswith(str((DATA / problem).resolve())):
        raise HTTPException(404)
    return FileResponse(p)


@router.post("/api/wb/run")
def start(body: dict = Body(...)):
    r = W().start(body["agent"], body.get("instructions", ""))
    return r.summary()


@router.get("/api/wb/runs")
def runs():
    return [r.summary() for r in W().runs.values()]


@router.get("/api/wb/run/{run_id}")
def run(run_id: str):
    r = W().runs.get(run_id)
    if not r:
        raise HTTPException(404)
    return r.summary(full=True)


@router.get("/api/wb/artifact/{run_id}/{aid}")
def artifact(run_id: str, aid: str):
    r = W().runs.get(run_id)
    if not r or aid not in r.artifacts:
        raise HTTPException(404)
    return r.artifacts[aid]


@router.get("/wb/file/{run_id}/{name}")
def file(run_id: str, name: str):
    p = OUT / run_id / name
    if not p.exists():
        raise HTTPException(404)
    return FileResponse(p, filename=name)


@router.get("/api/wb/audit/{run_id}")
def audit(run_id: str):
    return W().audit_for(run_id)


@router.get("/api/wb/verify")
def verify():
    return W().verify()


@router.post("/api/wb/item/{item_id}/challenge")
def challenge(item_id: str, body: dict = Body(...)):
    return W().challenge(item_id, body.get("question") or "Why?")


@router.post("/api/wb/item/{item_id}/{decision}")
def decide(item_id: str, decision: str, body: dict = Body(default={})):
    from jarvis.work.gateway import ValidationError
    
    w = W()
    try:
        if decision == "self-approve":
            transition(w.db, item_id, "approve", actor("jarvis"), reason="(agent attempting self-approval)")
            return {"denied": False}
        res = w.decide(item_id, "approve" if decision == "approve" else "reject", body.get("reason", ""))
        return res
    except ValidationError as e:
        raise HTTPException(400, f"Entry validation failed: {str(e)}")
    except Denied as e:
        r = w.runs.get(get_item(w.db, item_id)["proposal"].get("run_id"))
        if r:
            r.log("deny", f"DENIED in code: {e.reason} (attempted '{e.action}' on {item_id}) — logged, nothing changed")
        return JSONResponse({"denied": True, "reason": e.reason}, status_code=403)


@router.post("/api/wb/speed")
def speed(value: float = Body(..., embed=True)):
    WB["speed"] = value
    W().speed = value
    return {"speed": value}


@router.post("/api/wb/reset")
def reset():
    WB["wb"] = Workbench(fresh=True, speed=WB["speed"])
    events.publish("wb", kind="reset")
    return {"reset": True}
