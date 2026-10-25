"""The five workbench agents run end to end on the messy raw data, log every stage, and never write without approval."""
import time

import pytest
from openpyxl import load_workbook

from jarvis.work.roles import actor
from jarvis.work.transition import Denied, get_item, transition
from jarvis.workbench.core import Workbench


@pytest.fixture(scope="module")
def wb():
    return Workbench(fresh=True, speed=0)


def run(wb, agent):
    r = wb.start(agent)
    t0 = time.time()
    while r.status in ("queued", "running"):
        assert time.time() - t0 < 60, f"{agent} hung"
        time.sleep(0.05)
    return r


@pytest.mark.parametrize("agent,min_props", [("p1_timing", 4), ("p2_rollforward", 1), ("p3_bank", 6),
                                             ("p4_evidence", 1), ("p5_vendors", 6)])
def test_agent_runs_every_stage_and_waits_for_a_human(wb, agent, min_props):
    r = run(wb, agent)
    assert r.status == "awaiting_approval"
    assert all(s["status"] in ("done", "flag", "waiting") for s in r.stages)
    assert all(s["logs"] >= 2 for s in r.stages), "every stage logs its start and end"
    assert len(r.proposals) >= min_props
    assert all(get_item(wb.db, p)["state"] == "PROPOSED" for p in r.proposals), "nothing applied without approval"
    assert all(e["hash"] and e["audit_seq"] for e in r.logs), "every log line is an audit row"
    assert wb.verify()["ok"]


def test_cleaning_is_visible(wb):
    r = next(x for x in wb.runs.values() if x.agent == "p1_timing")
    fixes = [e for e in r.logs if e["level"] == "fix"]
    assert len(fixes) >= 5
    assert any(a["kind"] == "diff" for a in r.artifacts.values())


def test_agent_cannot_approve_its_own_work_and_the_attempt_is_logged(wb):
    r = next(x for x in wb.runs.values() if x.agent == "p1_timing")
    item = r.proposals[0]
    with pytest.raises(Denied):
        transition(wb.db, item, "approve", actor("jarvis"), reason="self-approval attempt")
    denied = wb.db.q("SELECT * FROM audit_log WHERE subject=? AND allowed=0", (item,))
    assert denied
    assert wb.verify()["ok"]


def test_controller_approval_goes_through_the_gateway(wb):
    r = next(x for x in wb.runs.values() if x.agent == "p1_timing")
    res = wb.decide(r.proposals[0], "approve", "ok")
    assert res["state"] == "POSTED"
    assert wb.db.one("SELECT 1 FROM ledger_entries WHERE work_item_id=?", (r.proposals[0],))


def test_package_formulas_and_late_change_supersedes_v1(wb):
    v1 = next(x for x in wb.runs.values() if x.agent == "p2_rollforward")
    path = next(a["payload"]["path"] for a in v1.artifacts.values() if a["kind"] == "file")
    ws = load_workbook(path)["IS"]
    ni = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "NET INCOME")
    assert str(ws.cell(ni, 11).value).startswith("="), "net income is a live formula"
    new = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "6455")
    assert ws.cell(new, 12).value == f"=SUM(C{new}:K{new})"
    v2 = run(wb, "p2_late")
    assert get_item(wb.db, v1.proposals[0])["state"] == "DISMISSED"
    assert get_item(wb.db, v2.proposals[0])["state"] == "PROPOSED"
    assert wb.verify()["ok"]


@pytest.mark.parametrize("agent", ["p1_timing", "p2_rollforward", "p3_bank", "p4_evidence", "p5_vendors", "p2_late"])
def test_every_run_is_scored_against_the_answer_key_and_beats_the_baseline(wb, agent):
    r = next((x for x in wb.runs.values() if x.agent == agent), None) or run(wb, agent)
    assert r.proof and r.proof["jarvis"] == r.proof["total"], [x for x in r.proof["rows"] if not x["jarvis_pass"]]
    assert r.proof["baseline"] < r.proof["total"]
    assert r.moments, "every run has at least one moment (caught / recovered / asked)"


def test_agent_defends_a_challenged_decision_from_logged_evidence(wb):
    r = next(x for x in wb.runs.values() if x.agent == "p3_bank")
    item = next(p for p in r.proposals if "DCE4343478" in get_item(wb.db, p)["title"])
    res = wb.challenge(item, "Why not just expense it all like we always do with PayPal? It's under materiality.")
    text = " ".join(res["points"])
    assert "P-0317" in text and "Prepaid" in text and "Log #" in text
    assert any(e.get("moment") == "defended" for e in r.logs)
    assert get_item(wb.db, item)["state"] == "PROPOSED", "defending does not approve anything"
    assert wb.verify()["ok"]
